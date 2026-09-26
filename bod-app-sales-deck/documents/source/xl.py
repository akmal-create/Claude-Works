from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
A='Arial'
def F(**k): return Font(name=A, **k)
INK='0B0D0C'; LIME='B8E62E'; GREY='6B7773'
YEL=PatternFill('solid',fgColor='FFFF00'); LIMEF=PatternFill('solid',fgColor=LIME); DARK=PatternFill('solid',fgColor=INK); SOFT=PatternFill('solid',fgColor='F3F5F4')
thin=Side(style='thin',color='E1E6E4'); B=Border(bottom=thin)
RS='"₹"#,##0;("₹"#,##0);"-"'
# ---------------- CALCULATOR ----------------
wb=Workbook(); ws=wb.active; ws.title='Calculator'
ws.sheet_view.showGridLines=False
for c,w in zip('ABCDEFG',[46,16,62,3,18,14,3]): ws.column_dimensions[c].width=w
ws['A1']='Bod App · Time-saved calculator'; ws['A1'].font=F(bold=True,size=16)
ws['A2']='Fill in the yellow cells with the agency\'s own numbers. Everything else updates on its own.'; ws['A2'].font=F(size=10,color=GREY)
ws['A3']='FIRST DRAFT · FOR REVIEW'; ws['A3'].font=F(size=8,bold=True,color='C23B22')
ws['A4']='THEIR TEAM'; ws['A4'].font=F(bold=True,size=9,color=GREY)
inputs=[('Team size (people)',10,'0','Example value. Everyone who will use Bod App.'),
('Average cost of one person, per hour (₹)',400,RS,'Example value. Salary plus overheads, divided by working hours.'),
('Minutes a day each person spends on admin',30,'0','Status updates, chasing, standups, working out what to bill. 30 min is the illustration used in the deck.'),
('Working days per month',22,'0','Standard working month.'),
('Share of that admin Bod App takes off the team',0.5,'0%','Assumption for the conversation, not a measured figure. Adjust with the client.'),
('Bod App plan price, per person per month (₹)',299,RS,'Pick 199 (Basic), 299 (Advanced) or 499 (Premium).')]
r=5
for lab,val,fmt,note in inputs:
    ws.cell(r,1,lab).font=F(size=10); c=ws.cell(r,2,val); c.font=F(size=10,color='0000FF',bold=True); c.fill=YEL; c.number_format=fmt; c.alignment=Alignment(horizontal='right')
    ws.cell(r,3,note).font=F(size=9,color=GREY); ws.cell(r,3).alignment=Alignment(wrap_text=True,vertical='center')
    for col in (1,2,3): ws.cell(r,col).border=B
    ws.row_dimensions[r].height=28; r+=1
dv=DataValidation(type='list',formula1='"199,299,499"',allow_blank=False); ws.add_data_validation(dv); dv.add('B10')
ws['A12']='WHAT IT MEANS EACH MONTH'; ws['A12'].font=F(bold=True,size=9,color=GREY)
outs=[('Hours spent on admin','=B5*B7/60*B8','#,##0','Team size × minutes a day ÷ 60 × working days.'),
('Cost of that admin','=B13*B6',RS,'Admin hours × cost per hour.'),
('Hours given back to real work','=B13*B9','#,##0','Admin hours × share Bod App takes off.'),
('Value of the hours given back','=B15*B6',RS,'Hours given back × cost per hour.'),
('Bod App cost','=B5*B10',RS,'Team size × plan price.'),
('Net gain','=B16-B17',RS,'Value of hours given back − Bod App cost.'),
('Return on spend','=IFERROR(B16/B17,0)','0.0"x"','How many rupees of time each rupee of Bod App gives back.'),
('Hours given back, per person','=IFERROR(B15/B5,0)','0.0','Hours given back ÷ team size.')]
r=13
for lab,f,fmt,note in outs:
    ws.cell(r,1,lab).font=F(size=10,bold=(lab in ('Net gain','Return on spend')))
    c=ws.cell(r,2,f); c.font=F(size=10,bold=True); c.number_format=fmt; c.alignment=Alignment(horizontal='right')
    if lab in ('Net gain','Return on spend'): 
        for col in (1,2): ws.cell(r,col).fill=LIMEF
    ws.cell(r,3,note).font=F(size=9,color=GREY)
    for col in (1,2,3): ws.cell(r,col).border=B
    ws.row_dimensions[r].height=22; r+=1
ws['E4']='PLANS'; ws['E4'].font=F(bold=True,size=9,color=GREY)
for i,(n,p) in enumerate([('Basic',199),('Advanced',299),('Premium',499)]):
    ws.cell(5+i,5,n).font=F(size=10); c=ws.cell(5+i,6,p); c.number_format=RS; c.font=F(size=10)
ws['E8']='Per person, per month'; ws['E8'].font=F(size=8,color=GREY)
ws['A23']='Notes'; ws['A23'].font=F(bold=True,size=10)
ws['A24']='The minutes-a-day and share-taken-off figures are conversation starters, not measured results. Use the client\'s own numbers wherever you can.'; ws['A24'].font=F(size=9,color=GREY)
ws.merge_cells('A24:C24'); ws['A24'].alignment=Alignment(wrap_text=True); ws.row_dimensions[24].height=28
ws['B9'].comment=Comment('Assumption: not a measured figure. Agree it with the client.','Bod Studio')
wb.save('05-Time-Saved-Calculator.xlsx')

# ---------------- QUOTE ----------------
wb=Workbook(); ws=wb.active; ws.title='Quote'; ws.sheet_view.showGridLines=False
for c,w in zip('ABCDEFG',[30,40,10,10,20,12,18]): ws.column_dimensions[c].width=w
ws['A1']='Bod App · Quote'; ws['A1'].font=F(bold=True,size=16)
ws['A2']='Bod Studio (Storibod Creatives) · Kochi, Kerala · bodstudio.com · +91 7356 333 965 · info@storibodcreatives.com'; ws['A2'].font=F(size=9,color=GREY)
ws['A3']='FIRST DRAFT · FOR REVIEW'; ws['A3'].font=F(size=8,bold=True,color='C23B22')
meta=[('Prepared for','[Agency name]'),('Contact','[Name, phone, email]'),('Quote number','BOD-2026-001'),('Date','=TODAY()'),('Valid until','=B8+15')]
for i,(k,v) in enumerate(meta):
    ws.cell(5+i,1,k).font=F(size=10,color=GREY); c=ws.cell(5+i,2,v); c.font=F(size=10,bold=True,color='0000FF' if not str(v).startswith('=') or k=='Date' else '000000')
    if k in ('Prepared for','Contact','Quote number','Date'): c.fill=YEL
    if k in ('Date','Valid until'): c.number_format='DD MMM YYYY'; c.alignment=Alignment(horizontal='left')
hdr=['Item','What\'s included','People','Months','Price / person / month','Discount','Amount']
for j,h in enumerate(hdr):
    c=ws.cell(11,1+j,h); c.font=F(bold=True,size=9,color='FFFFFF'); c.fill=DARK; c.alignment=Alignment(horizontal='left' if j<2 else 'right',vertical='center')
ws.row_dimensions[11].height=22
lines=[('Free month','Basic setup, up to 3 people',3,1,0,0),('Advanced plan','Everything we use at Storibod',8,12,299,0.1),('','',None,None,None,None),('','',None,None,None,None),('','',None,None,None,None)]
for i,(a,b,p,m,pr,dsc) in enumerate(lines):
    r=12+i
    for j,v in enumerate([a,b,p,m,pr,dsc]):
        c=ws.cell(r,1+j,v); c.font=F(size=10,color='0000FF'); c.fill=YEL; c.border=B
    ws.cell(r,5).number_format=RS; ws.cell(r,6).number_format='0%'
    for col in (3,4,5,6): ws.cell(r,col).alignment=Alignment(horizontal='right')
    g=ws.cell(r,7,f'=IF(OR(C{r}="",D{r}="",E{r}=""),0,C{r}*D{r}*E{r}*(1-F{r}))'); g.number_format=RS; g.font=F(size=10,bold=True); g.border=B
ws['F18']='Subtotal'; ws['G18']='=SUM(G12:G16)'
ws['F19']='GST rate'; ws['G19']=0.18; ws['G19'].fill=YEL; ws['G19'].font=F(size=10,color='0000FF'); ws['G19'].number_format='0%'
ws['G19'].comment=Comment('Assumption: 18% GST on SaaS in India. Confirm with your accountant before sending.','Bod Studio')
ws['F20']='GST'; ws['G20']='=G18*G19'
ws['F21']='Total'; ws['G21']='=G18+G20'
for r in (18,20,21): ws.cell(r,7).number_format=RS
for r in (18,19,20,21):
    ws.cell(r,6).font=F(size=10,bold=(r==21)); ws.cell(r,6).alignment=Alignment(horizontal='right')
    if r!=19: ws.cell(r,7).font=F(size=10 if r!=21 else 12,bold=True)
    ws.cell(r,7).alignment=Alignment(horizontal='right')
ws['F21'].fill=LIMEF; ws['G21'].fill=LIMEF
ws['A23']='How to use'; ws['A23'].font=F(bold=True,size=10)
notes=['Yellow cells are for you to edit. The two filled rows are examples: replace or clear them.','Plan prices: Basic ₹199, Advanced ₹299, Premium ₹499 per person, per month.','Terms: [add payment terms, billing cycle and cancellation terms once confirmed].']
for i,t in enumerate(notes):
    ws.cell(24+i,1,t).font=F(size=9,color=GREY); ws.merge_cells(start_row=24+i,start_column=1,end_row=24+i,end_column=7)
wb.save('04-Quote-Template.xlsx')
print('saved')
